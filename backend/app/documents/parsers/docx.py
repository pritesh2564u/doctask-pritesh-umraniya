from pathlib import Path

from docx import Document

from app.documents.models import ParsedBlock, ParsedDocument, SourceLocation
from app.documents.parser import DocumentParser


class DOCXParser(DocumentParser):

    def parse(self, file_path: Path) -> ParsedDocument:
        document = Document(str(file_path))

        blocks: list[ParsedBlock] = []

        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()

            if not text:
                continue

            blocks.append(
                ParsedBlock(
                    text=text,
                    location=SourceLocation(
                        paragraph=index,
                    ),
                )
            )

        return ParsedDocument(
            filename=file_path.name,
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            blocks=blocks,
        )